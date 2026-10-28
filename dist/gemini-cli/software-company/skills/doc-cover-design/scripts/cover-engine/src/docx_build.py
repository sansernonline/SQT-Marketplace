"""Build docx/<p>-srs-cover-sample.docx with brandkit: page 1 = cover, page 2 = a short normal body page.

Cover recipe
- background: bg/<p>-srs-bg.png (2480x3508, 300 dpi) via brandkit.anchor_picture() on the FIRST body paragraph,
  page-relative, 0/0/21.0/29.7 cm, behindDoc. Never in the header (Word fades header pictures).
- every live text block from build/layout/<p>-srs.json becomes one paragraph with an absolute frame
  (w:framePr, page-anchored, x/y/w in twips = px * 9.6) and exact line spacing = CSS line-height.
  Font size pt = px * 0.48 (150 dpi). Each weight is its own installed family (see fonts.py).
- build/calib.json (written by calibrate.py) holds per-style y/x corrections measured from Word's own render.
- page 2 starts with pageBreakBefore on its heading, so no empty break paragraph can spill into a blank page.
"""
import json
import os
import sys
from pathlib import Path

from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
# brandkit.py: BRANDKIT_DIR, else branded-document-design in the superuser plugin (sibling plugin folder),
# else a sibling skill in this plugin (older layout)
SKILLS = Path(__file__).resolve().parents[4]
_candidates = [SKILLS.parent.parent / "superuser" / "skills" / "branded-document-design" / "scripts",
               SKILLS / "branded-document-design" / "scripts"]
_default = next((c for c in _candidates if c.exists()), _candidates[0])
sys.path.insert(0, os.environ.get("BRANDKIT_DIR", str(_default)))
import brandkit as bk  # noqa: E402
from fonts import word_family  # noqa: E402

TW = 9.6          # twips per px at 150 dpi (1440 / 150)
PT = 0.48         # pt per px at 150 dpi (72 / 150)
Z_BG = 251658240

PAGE2 = {   # body fonts, brand tokens and a short real-data body page per project
    "litenet": dict(thai="IBM Plex Sans Thai", latin="IBM Plex Sans Thai",
                    footer="หน้า", brand=dict(brand="3F7A00", brand_deep="0A0A0A", brand_tint="EEF5E6"),
                    h1="1. เกี่ยวกับเอกสารนี้",
                    para="LiteNet คือเครือข่ายส่วนตัวแบบ mesh ที่ตั้งบนเซิร์ฟเวอร์ของคุณเอง หน้านี้เป็นหน้าเนื้อหาตัวอย่าง "
                         "ใช้ตรวจว่าปกจบที่หน้า 1 และเนื้อหาเริ่มที่หน้า 2 พอดี",
                    head=["รายการ", "ค่า"],
                    rows=[["รหัสเอกสาร", "LITENET-SRS"], ["เวอร์ชัน", "0.3"], ["สถานะ", "DRAFT"], ["วันที่", "2026-10-15"]]),
    "swipeos": dict(thai="Anuphan", latin="Space Grotesk",
                    footer="หน้า", brand=dict(brand="0B7A3E", brand_deep="08170E", brand_tint="E8F2EC"),
                    h1="1. เกี่ยวกับเอกสารนี้",
                    para="SwipeOS เปลี่ยนท่ามือหน้าเว็บแคมเป็นคีย์ลัด ทำงานออฟไลน์ ไม่ส่งภาพออกนอกเครื่อง "
                         "หน้านี้เป็นหน้าเนื้อหาตัวอย่าง ใช้ตรวจว่าปกจบที่หน้า 1 และเนื้อหาเริ่มที่หน้า 2 พอดี",
                    head=["รายการ", "ค่า"],
                    rows=[["รหัสเอกสาร", "SWIPEOS-SRS"], ["เวอร์ชัน", "0.3"], ["สถานะ", "DRAFT"], ["วันที่", "2026-10-13"]]),
    "telex": dict(thai="Space Grotesk", latin="Space Grotesk",
                  footer="Page", brand=dict(brand="6D28D9", brand_deep="14121F", brand_tint="F1EEFC"),
                  h1="1. About this document",
                  para="Telex puts SSH, Telnet, RDP, VNC, SFTP, serial and X11 in one Windows app. This is a sample "
                       "body page: it shows the cover ends on page 1 and the body starts on page 2.",
                  head=["Item", "Value"],
                  rows=[["Document", "TLX-SRS"], ["Version", "0.3"], ["Status", "DRAFT"], ["Date", "2026-10-11"]]),
}


def style_key(b):
    r = next(r for r in b["runs"] if "text" in r)
    return f'{r["family"]}|{r["weight"]}|{r["sizePx"]}|{b["lineHeightPx"]}'


def frame_paragraph(doc, b, calib):
    dx, dy = calib.get(style_key(b), [0, 0])
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    ppr = p._p.get_or_add_pPr()
    spacing = ppr.find(qn("w:spacing"))
    spacing.set(qn("w:line"), str(round(b["lineHeightPx"] * TW)))
    spacing.set(qn("w:lineRule"), "exact")
    x, y = (b["x"] - dx) * TW, (b["y"] - dy) * TW
    w = min(b["w"] + 80, 1240 - b["x"]) * TW            # slack so Word's metrics never force an extra wrap
    ppr.insert(0, parse_xml(f'<w:framePr {nsdecls("w")} w:w="{round(w)}" w:wrap="around" w:hAnchor="page" '
                            f'w:vAnchor="page" w:x="{round(x)}" w:y="{round(y)}"/>'))
    for r in b["runs"]:
        if r.get("br"):
            p.add_run().add_break()
            continue
        name, bold = word_family(r["family"], r["weight"])
        run = bk.style_run(p.add_run(r["text"]), size=round(r["sizePx"] * PT * 2) / 2, color=r["color"].lstrip("#"),
                           bold=bold, font=name)
        if r["letterSpacingPx"]:
            rpr = run._r.get_or_add_rPr()
            rpr.append(bk._el("w:spacing", val=round(r["letterSpacingPx"] * PT * 20)))
            bk.order_rpr(rpr)
    # Thai segments: brandkit thai_runs puts every Thai run on bk.FONT, so point it at this block's own font
    first = next(r for r in b["runs"] if "text" in r)
    bk.FONT = word_family(first["family"], first["weight"])[0]
    bk.thai_runs(p._p)
    return p


def no_compress(settings):
    """<w:doNotAutoCompressPictures/> in its schema slot (Word rejects out-of-order settings children)."""
    after = ("forceUpgrade", "captions", "readModeInkLockDown", "smartTagType", "schemaLibrary",
             "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator")
    el = bk._el("w:doNotAutoCompressPictures")
    nxt = next((c for c in settings if c.tag in {qn("w:" + t) for t in after}), None)
    if nxt is None:
        settings.append(el)
    else:
        nxt.addprevious(el)


def build(project, calib):
    cfg = PAGE2[project]
    lay = json.loads((ROOT / "build" / "layout" / f"{project}-srs.json").read_text(encoding="utf-8"))
    bk.FONT, bk.LATIN_FONT = cfg["thai"], cfg["latin"]
    bk.use_brand(**cfg["brand"])
    d = bk.BrandDoc(footer_text=cfg["footer"])
    doc = d.doc
    doc.sections[0].different_first_page_header_footer = True     # no footer on the cover
    # 1) background anchored to the first paragraph (not the header)
    first = doc.add_paragraph()
    first.paragraph_format.space_before = first.paragraph_format.space_after = Pt(0)
    first.paragraph_format.line_spacing = Pt(1)
    bk.anchor_picture(first, str(ROOT / "bg" / f"{project}-srs-bg.png"), 0, 0, 21.0, 29.7, Z_BG)
    # 2) live text, framed at the preview's positions
    for b in lay["blocks"]:
        frame_paragraph(doc, b, calib.get(project, {}))
    # 3) body page
    bk.FONT, bk.LATIN_FONT = cfg["thai"], cfg["latin"]
    h = d.h1(cfg["h1"], space_before=0)
    h.paragraph_format.page_break_before = True
    d.para(cfg["para"])
    d.table(cfg["head"], cfg["rows"], widths=[3000, 6638])
    for el in list(doc.element.body)[list(doc.element.body).index(h._p):]:   # body page only, not the cover frames
        bk.thai_runs(el)
    bk.thai_runs(doc.sections[0].footer._element)
    # keep the 300 dpi background: Word must not downsample pictures when it re-saves (word_pdf.py)
    no_compress(doc.settings.element)
    out = ROOT / "docx" / f"{project}-srs-cover-sample.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(out))
    return out


def main():
    cpath = ROOT / "build" / "calib.json"
    calib = json.loads(cpath.read_text(encoding="utf-8")) if cpath.exists() else {}
    return [build(p, calib) for p in PAGE2]


def self_test():
    """python docx_build.py --self-test: one framed text block and the settings slot; no layout json, background or Word needed."""
    block = {"x": 100, "y": 200, "w": 600, "lineHeightPx": 50,
             "runs": [{"text": "Hello ไทย", "family": "Space Grotesk", "weight": 600, "sizePx": 40,
                       "color": "#112233", "letterSpacingPx": 0.5}]}
    assert style_key(block) == "Space Grotesk|600|40|50"
    d = bk.BrandDoc()
    p = frame_paragraph(d.doc, block, {"Space Grotesk|600|40|50": [10, 20]})
    ppr = p._p.pPr
    frame = ppr.find(qn("w:framePr"))
    assert frame.get(qn("w:x")) == str(round((100 - 10) * TW)) and frame.get(qn("w:y")) == str(round((200 - 20) * TW))
    assert frame.get(qn("w:w")) == str(round(680 * TW)) and frame.get(qn("w:hAnchor")) == "page"
    spacing = ppr.find(qn("w:spacing"))
    assert spacing.get(qn("w:line")) == "480" and spacing.get(qn("w:lineRule")) == "exact"
    assert [r.text for r in p.runs] == ["Hello ", "ไทย"], [r.text for r in p.runs]
    assert p.runs[0].font.size == Pt(19.0) and p.runs[0].font.name == "Space Grotesk SemiBold"
    settings = bk._el("w:settings")
    settings.append(bk._el("w:decimalSymbol", val="."))
    no_compress(settings)
    assert [c.tag for c in settings][0] == qn("w:doNotAutoCompressPictures")    # before decimalSymbol, schema order
    print("self-test ok")


if __name__ == "__main__":
    if sys.argv[1:] == ["--self-test"]:
        self_test()
        sys.exit()
    for f in main():
        print("docx", f.name)
