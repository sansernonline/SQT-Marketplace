"""pptx_to_html.py — แปลง .pptx เป็นเด็ค .html ไฟล์เดียว ที่ทุกชิ้นเลื่อนเข้ามาทีละกลุ่ม

    python pptx_to_html.py deck.pptx -o deck.html --accent E4572E --title "ชื่อเด็ค"

ได้: ไฟล์เดียว (รูปฝังในไฟล์) · ตำแหน่งตรงกับ PowerPoint · โน้ตผู้พูดจาก notes · สารบัญจาก section
animation ใส่ให้เอง (หัวข้อพุ่งจากซ้าย · ข้อความลอยขึ้น · กล่องเด้ง · เส้นลากตามทิศ · ตารางขึ้นทีละแถว · กราฟแท่งโต · ตัวเลขใหญ่นับขึ้น)
animation แบบ "On Click" ใน PowerPoint กลายเป็นขั้นในหน้า (→ เปิดทีละขั้นก่อนไปหน้าถัดไป)
ของที่แปลงไม่ได้ (SmartArt · วิดีโอ · กราฟบางชนิด · รูปทรงแปลก) ขึ้นเป็นกล่องลายเตือน และพิมพ์รายการใน stdout
ต้องมี python-pptx
"""
from __future__ import annotations

import argparse
import base64
import colorsys
import html
import math
import re
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Emu

ASSETS = Path(__file__).resolve().parent.parent / "assets"
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
      "p14": "http://schemas.microsoft.com/office/powerpoint/2010/main"}
SCHEME_ALIAS = {"bg1": "lt1", "tx1": "dk1", "bg2": "lt2", "tx2": "dk2"}
FONT_STACK = "Tahoma,'Leelawadee UI',Loma,sans-serif"
warnings: list[str] = []


def q(el, path):
    return el.find(path, NS) if el is not None else None


def esc(s: str) -> str:
    return html.escape(s, quote=True)


class Deck:
    def __init__(self, path: Path):
        self.prs = Presentation(str(path))
        self.k = 1920 / self.prs.slide_width            # px ต่อ EMU · กว้าง 1920 เสมอ
        self.W, self.H = 1920, round(self.prs.slide_height * self.k)
        theme = self._theme()
        self.scheme = {}
        for c in q(theme, ".//a:clrScheme") if theme is not None else []:
            v = c[0]
            self.scheme[etree.QName(c).localname] = v.get("lastClr", "000000") if etree.QName(v).localname == "sysClr" else v.get("val")
        self.fonts = {"+mj": self._face(theme, "majorFont"), "+mn": self._face(theme, "minorFont")}
        self.markers: set[str] = set()

    def _theme(self):
        master = self.prs.slide_masters[0]
        for rel in master.part.rels.values():
            if rel.reltype.endswith("/theme"):
                return etree.fromstring(rel.target_part.blob)
        return None

    @staticmethod
    def _face(theme, kind):
        el = q(theme, f".//a:{kind}/a:latin")
        return el.get("typeface") if el is not None else ""

    def px(self, emu) -> float:
        return round((emu or 0) * self.k, 1)

    # ---------------------------------------------------------------- colour
    def color(self, holder) -> str | None:
        """holder = element ที่มี a:solidFill (หรือเป็น solidFill เอง) · คืน '#RRGGBB'"""
        if holder is None:
            return None
        fill = holder if etree.QName(holder).localname == "solidFill" else q(holder, "a:solidFill")
        if fill is None or not len(fill):
            return None
        return self.clr(fill[0])

    def clr(self, c) -> str | None:
        name = etree.QName(c).localname
        if name == "srgbClr":
            hexv = c.get("val")
        elif name == "schemeClr":
            hexv = self.scheme.get(SCHEME_ALIAS.get(c.get("val"), c.get("val")), "000000")
        elif name == "sysClr":
            hexv = c.get("lastClr", "000000")
        elif name == "prstClr":
            hexv = {"black": "000000", "white": "FFFFFF", "red": "FF0000"}.get(c.get("val"), "000000")
        else:
            return None
        r, g, b = (int(hexv[i:i + 2], 16) / 255 for i in (0, 2, 4))
        mod, off = c.find("a:lumMod", NS), c.find("a:lumOff", NS)
        if mod is not None or off is not None:
            h, l, s = colorsys.rgb_to_hls(r, g, b)
            l = l * (int(mod.get("val")) / 100000 if mod is not None else 1) + (int(off.get("val")) / 100000 if off is not None else 0)
            r, g, b = colorsys.hls_to_rgb(h, min(1, max(0, l)), s)
        tint, shade = c.find("a:tint", NS), c.find("a:shade", NS)
        if tint is not None:                              # ผสมขาว: ยิ่งค่าน้อยยิ่งอ่อน
            t = int(tint.get("val")) / 100000
            r, g, b = (x * t + (1 - t) for x in (r, g, b))
        if shade is not None:                             # ผสมดำ
            t = int(shade.get("val")) / 100000
            r, g, b = (x * t for x in (r, g, b))
        out = "#%02X%02X%02X" % (round(r * 255), round(g * 255), round(b * 255))
        alpha = c.find("a:alpha", NS)
        if alpha is not None:
            out += "%02X" % round(int(alpha.get("val")) / 100000 * 255)
        return out


# ---------------------------------------------------------------- text
def is_title(shape) -> bool:
    """ชื่อเรื่องจริง (TITLE · CENTER_TITLE) · SUBTITLE ไม่นับ"""
    return shape.is_placeholder and str(shape.placeholder_format.type).split(" ")[0] in ("TITLE", "CENTER_TITLE")


def layout_placeholders(shape, slide):
    """placeholder ของ layout และ master ที่ shape นี้รับค่ามา"""
    if not shape.is_placeholder:
        return []
    idx, typ = shape.placeholder_format.idx, shape.placeholder_format.type
    out = []
    for owner in (slide.slide_layout, slide.slide_layout.slide_master):
        for ph in owner.placeholders:
            if ph.placeholder_format.idx == idx or ph.placeholder_format.type == typ:
                out.append(ph)
                break
    return out


def inherited(deck, shape, slide, lvl: int, attr: str, sub: str = "/a:defRPr"):
    """หาค่าที่ไม่ได้ตั้งไว้ตรง ๆ ตามลำดับที่ PowerPoint ใช้: lstStyle ของ shape → placeholder ของ layout → master → txStyles
    sub="/a:defRPr" = ค่าตัวอักษร (sz · b · i · fill) · sub="" = ค่าย่อหน้า (algn)"""
    tag = f"a:lvl{lvl + 1}pPr"
    chain = [q(shape._element, f".//a:lstStyle/{tag}{sub}")]
    chain += [q(ph._element, f".//a:lstStyle/{tag}{sub}") for ph in layout_placeholders(shape, slide)]
    styles = q(slide.slide_layout.slide_master._element, "p:txStyles")
    if shape.is_placeholder:
        kind = "p:titleStyle" if is_title(shape) else "p:bodyStyle"
        chain.append(q(styles, f"{kind}/{tag}{sub}"))
    else:
        chain.append(q(styles, f"p:otherStyle/{tag}{sub}"))
    for el in chain:
        if el is not None and (attr == "fill" and q(el, "a:solidFill") is not None or el.get(attr) is not None):
            return el
    return None


def run_css(deck, shape, slide, para, rpr, scale, fallback=None):
    lvl = int(para.get("lvl", 0)) if para is not None else 0
    def get(attr):
        if rpr is not None and (rpr.get(attr) is not None if attr != "fill" else q(rpr, "a:solidFill") is not None):
            return rpr
        ppr_def = q(para, "a:pPr/a:defRPr")
        if ppr_def is not None and (ppr_def.get(attr) is not None if attr != "fill" else q(ppr_def, "a:solidFill") is not None):
            return ppr_def
        return inherited(deck, shape, slide, lvl, attr)
    css = []
    sz = get("sz")
    size = (int(sz.get("sz")) / 100 if sz is not None else 18) * scale
    css.append(f"font-size:{size * 2:.1f}px")                 # 1 pt = 2 px บนเวที 1920 (13.33 นิ้ว)
    b = get("b")
    if b is not None and b.get("b") in ("1", "true"):
        css.append("font-weight:700")
    it = get("i")
    if it is not None and it.get("i") in ("1", "true"):
        css.append("font-style:italic")
    col = get("fill")
    explicit = col is not None and (col is rpr or col is q(para, "a:pPr/a:defRPr"))
    if not explicit and fallback is None:
        ref = q(shape._element, "p:style/a:fontRef")
        fallback = deck.clr(ref[0]) if ref is not None and len(ref) else None
    if explicit or (col is not None and fallback is None):
        css.append(f"color:{deck.color(col)}")
    elif fallback:
        css.append(f"color:{fallback}")
    u = rpr.get("u") if rpr is not None else None
    if u and u != "none":
        css.append("text-decoration:underline")
    face = None
    for kind in ("a:cs", "a:latin"):
        el = q(rpr, kind)
        if el is not None and el.get("typeface"):
            face = el.get("typeface")
            break
    if face:
        face = deck.fonts.get(face[:3], face) if face.startswith("+") else face
        if face:
            css.append(f"font-family:'{face}',{FONT_STACK}")
    return ";".join(css), size


def text_html(deck, shape, slide) -> str:
    body = q(shape._element, ".//p:txBody") if shape._element.find(".//p:txBody", NS) is not None else q(shape._element, ".//a:txBody")
    if body is None:
        return ""
    bpr = q(body, "a:bodyPr")
    if bpr.get("vert", "horz") != "horz":
        warnings.append("ข้อความแนวตั้ง (แสดงเป็นแนวนอน)")
    auto = q(bpr, "a:normAutofit")
    scale = int(auto.get("fontScale", 100000)) / 100000 if auto is not None else 1
    ins = [deck.px(int(bpr.get(k, d))) for k, d in (("tIns", 45720), ("rIns", 91440), ("bIns", 45720), ("lIns", 91440))]
    anchors = [bpr] + [q(ph._element, ".//a:bodyPr") for ph in layout_placeholders(shape, slide)]
    anchor_val = next((b.get("anchor") for b in anchors if b is not None and b.get("anchor")), None)
    anchor = {"ctr": "center", "b": "flex-end"}.get(anchor_val, "flex-start")
    nvsp = q(shape._element, ".//p:cNvSpPr")
    textbox = nvsp is not None and nvsp.get("txBox") == "1"
    # ข้อความใน autoshape ที่ไม่ได้ตั้งแนวไว้ PowerPoint จัดกลาง · กล่องข้อความชิดซ้าย
    default_algn = "ctr" if not shape.is_placeholder and not textbox and etree.QName(shape._element).localname == "sp" else None
    wrap = "white-space:nowrap;" if bpr.get("wrap") == "none" else ""
    out = []
    for para in body.findall("a:p", NS):
        ppr = q(para, "a:pPr")
        style = []
        algn = ppr.get("algn") if ppr is not None else None
        if algn is None and default_algn:                     # autoshape: otherStyle ของ master ไม่มีผลกับแนว
            lst = q(shape._element, f".//a:lstStyle/a:lvl{int(ppr.get('lvl', 0)) + 1 if ppr is not None else 1}pPr")
            algn = lst.get("algn") if lst is not None and lst.get("algn") else default_algn
        if algn is None:
            src = inherited(deck, shape, slide, int(para.get("lvl", 0)) if ppr is None else int(ppr.get("lvl", 0)), "algn", "")
            algn = src.get("algn") if src is not None else default_algn
        if algn in ("ctr", "r", "just"):
            style.append("text-align:" + {"ctr": "center", "r": "right", "just": "justify"}[algn])
        for tag, prop in (("a:spcBef", "margin-top"), ("a:spcAft", "margin-bottom")):
            pts, pct = q(ppr, f"{tag}/a:spcPts"), q(ppr, f"{tag}/a:spcPct")
            if pts is not None:
                style.append(f"{prop}:{int(pts.get('val')) / 100 * 2:.1f}px")
            elif pct is not None:                              # % ของขนาดตัวอักษร
                style.append(f"{prop}:{int(pct.get('val')) / 100000:.2f}em")
        ln, ln_pts = q(ppr, "a:lnSpc/a:spcPct"), q(ppr, "a:lnSpc/a:spcPts")
        if ln is not None:
            style.append(f"line-height:{int(ln.get('val')) / 100000 * 1.2:.2f}")
        elif ln_pts is not None:
            style.append(f"line-height:{int(ln_pts.get('val')) / 100 * 2:.1f}px")
        bullet = ""
        bu = q(ppr, "a:buChar")
        if bu is not None:
            mar = deck.px(int(ppr.get("marL", 0)))
            style.append(f"padding-left:{mar}px")
            bullet = f'<span class="bu">{esc(bu.get("char"))}</span>'
        spans, size, biggest = [], 18, 0
        for r in para:
            name = etree.QName(r).localname
            if name in ("r", "fld"):
                css, size = run_css(deck, shape, slide, para, q(r, "a:rPr"), scale)
                biggest = max(biggest, size)
                t = q(r, "a:t")
                spans.append(f'<span style="{css}">{esc(t.text or "") if t is not None else ""}</span>')
            elif name == "br":
                spans.append("<br>")
        if not spans:
            css, size = run_css(deck, shape, slide, para, q(para, "a:endParaRPr"), scale)
            biggest = size
            spans.append(f'<span style="{css}">&#8203;</span>')
        # ความสูงบรรทัดตามตัวที่ใหญ่สุดในย่อหน้า แบบ PowerPoint (ไม่ใช่ 36px ของกล่อง)
        style.insert(0, f"font-size:{biggest * 2:.1f}px")
        out.append(f'<p style="{";".join(style)}">{bullet}{"".join(spans)}</p>')
    pad = " ".join(f"{v}px" for v in ins)
    return f'<div class="tx" style="{wrap}padding:{pad};justify-content:{anchor}">{"".join(out)}</div>'


# ---------------------------------------------------------------- shapes
def geometry(deck, el, off, shape=None):
    xfrm = q(el, "p:spPr/a:xfrm")
    if xfrm is None:
        xfrm = q(el, "p:xfrm")
    if xfrm is None:
        if shape is not None and shape.is_placeholder and shape.width:   # ตำแหน่งรับมาจาก layout
            x, y = off(shape.left, shape.top)
            return {"x": deck.px(x), "y": deck.px(y), "w": deck.px(shape.width * off.sx), "h": deck.px(shape.height * off.sy),
                    "rot": 0, "fh": False, "fv": False}
        return None
    o, e = q(xfrm, "a:off"), q(xfrm, "a:ext")
    x, y = off(int(o.get("x")), int(o.get("y")))
    w, h = int(e.get("cx")) * off.sx, int(e.get("cy")) * off.sy
    return {"x": deck.px(x), "y": deck.px(y), "w": deck.px(w), "h": deck.px(h),
            "rot": int(xfrm.get("rot", 0)) / 60000, "fh": xfrm.get("flipH") == "1", "fv": xfrm.get("flipV") == "1"}


class Offset:
    """แปลงพิกัดลูกในกลุ่ม (ซ้อนกันได้) เป็นพิกัดสไลด์"""
    def __init__(self, parent=None, grp=None):
        self.parent, self.fx, self.fy = parent, 1.0, 1.0
        if grp is not None:
            x = q(grp, "p:grpSpPr/a:xfrm")
            o, e, co, ce = (q(x, t) for t in ("a:off", "a:ext", "a:chOff", "a:chExt"))
            self.o = (int(o.get("x")), int(o.get("y")))
            self.c = (int(co.get("x")), int(co.get("y")))
            self.fx = int(e.get("cx")) / max(1, int(ce.get("cx")))
            self.fy = int(e.get("cy")) / max(1, int(ce.get("cy")))
        self.sx = self.fx * (parent.sx if parent else 1)
        self.sy = self.fy * (parent.sy if parent else 1)

    def __call__(self, x, y):
        if hasattr(self, "o"):
            x, y = self.o[0] + (x - self.c[0]) * self.fx, self.o[1] + (y - self.c[1]) * self.fy
        return self.parent(x, y) if self.parent else (x, y)


def box_style(g, extra=""):
    s = f"left:{g['x']}px;top:{g['y']}px;width:{g['w']}px;height:{g['h']}px"
    if g["rot"]:
        s += f";transform:rotate({g['rot']:.1f}deg)"
    return s + extra


def paint(deck, el):
    """(สีพื้น, สีเส้น, ความหนาเส้น px) · ไม่ได้ตั้งไว้ใน spPr → ใช้ p:style ของ shape (fillRef · lnRef) แบบที่ PowerPoint ทำ"""
    sppr, style = q(el, "p:spPr"), q(el, "p:style")
    declared = any(q(sppr, t) is not None for t in ("a:solidFill", "a:noFill", "a:gradFill", "a:blipFill", "a:pattFill"))
    fill = deck.color(q(sppr, "a:solidFill"))
    ref = q(style, "a:fillRef")
    if not declared and ref is not None and ref.get("idx") != "0" and len(ref):
        fill = deck.clr(ref[0])
    ln = q(sppr, "a:ln")
    stroke, width = None, 0
    if ln is None or q(ln, "a:noFill") is None:
        stroke = deck.color(ln) if ln is not None else None
        lref = q(style, "a:lnRef")
        if stroke is None and lref is not None and lref.get("idx") != "0" and len(lref):
            stroke = deck.clr(lref[0])
        width = deck.px(int(ln.get("w", 12700)) if ln is not None else 12700 * int(lref.get("idx", 1)) if lref is not None else 12700)
    return fill, stroke, width


def fill_line(deck, el):
    css = []
    fill, stroke, width = paint(deck, el)
    if fill:
        css.append(f"background:{fill}")
    grad = q(el, "p:spPr/a:gradFill")
    if grad is not None:
        stops = [deck.color(gs) for gs in grad.findall(".//a:gs", NS)]
        stops = [s for s in stops if s]
        if stops:
            css.append(f"background:linear-gradient(180deg,{','.join(stops)})")
    if stroke:
        css.append(f"border:{width:.1f}px solid {stroke}")
    return css


POLY = {"diamond": "50,0 100,50 50,100 0,50", "triangle": "50,0 100,100 0,100",
        "rightArrow": "0,25 60,25 60,0 100,50 60,100 60,75 0,75", "leftArrow": "100,25 40,25 40,0 0,50 40,100 40,75 100,75",
        "chevron": "0,0 75,0 100,50 75,100 0,100 25,50", "homePlate": "0,0 80,0 100,50 80,100 0,100",
        "hexagon": "25,0 75,0 100,50 75,100 25,100 0,50", "parallelogram": "25,0 100,0 75,100 0,100"}


def line_svg(deck, el, g, k):
    ln = q(el, "p:spPr/a:ln")
    col = paint(deck, el)[1] or "#808080"
    wpx = max(1, deck.px(int(ln.get("w", 12700)) if ln is not None else 12700))
    mid = "m" + col.lstrip("#")[:6]
    head = q(ln, "a:headEnd") if ln is not None else None
    tail = q(ln, "a:tailEnd") if ln is not None else None
    has = lambda e: e is not None and e.get("type", "none") != "none"
    if has(head) or has(tail):
        deck.markers.add(col[:7])
    x1, y1, x2, y2 = (g["w"], 0, 0, g["h"]) if g["fh"] else (0, 0, g["w"], g["h"])
    if g["fv"]:
        y1, y2 = y2, y1
    prst = (q(el, "p:spPr/a:prstGeom").get("prst") if q(el, "p:spPr/a:prstGeom") is not None else "line")
    if prst.startswith("bentConnector"):
        mx = (x1 + x2) / 2
        pts = f"{x1},{y1} {mx},{y1} {mx},{y2} {x2},{y2}"
    else:
        pts = f"{x1},{y1} {x2},{y2}"
    dx, dy = x2 - x1, y2 - y1
    direction = ("r" if dx >= 0 else "l") if abs(dx) >= abs(dy) else ("d" if dy >= 0 else "u")
    marks = (f' marker-start="url(#{mid})"' if has(head) else "") + (f' marker-end="url(#{mid})"' if has(tail) else "")
    dash = ' stroke-dasharray="8 6"' if q(ln, "a:prstDash") is not None and q(ln, "a:prstDash").get("val") != "solid" else ""
    w, h = max(g["w"], 1), max(g["h"], 1)
    return (f'<svg data-a="draw" class="s ln" data-dir="{direction}" style="--k:{k};{box_style(dict(g, w=w, h=h))}" viewBox="0 0 {w} {h}">'
            f'<polyline points="{pts}" fill="none" stroke="{col[:7]}" stroke-width="{wpx}"{dash}{marks}/></svg>')


def picture_html(deck, part, el, g, k, anim="pop"):
    blip = q(el, ".//a:blip")
    rid = blip.get(f"{{{NS['r']}}}embed") if blip is not None else None
    if not rid or rid not in part.rels:
        return unsupported(g, k, "รูปที่ลิงก์จากนอกไฟล์")
    img = part.related_part(rid)
    src = f"data:{img.content_type};base64,{base64.b64encode(img.blob).decode()}"
    crop = q(el, ".//a:srcRect")
    if crop is not None and any(crop.get(s) for s in "ltrb"):
        l, t, r, b = (int(crop.get(s, 0)) / 100000 for s in "ltrb")
        fw, fh = g["w"] / max(0.01, 1 - l - r), g["h"] / max(0.01, 1 - t - b)
        return (f'<div data-a="{anim}" class="s" style="--k:{k};{box_style(g)};overflow:hidden">'
                f'<img class="pic" style="position:absolute;left:{-l * fw:.1f}px;top:{-t * fh:.1f}px;width:{fw:.1f}px;height:{fh:.1f}px" src="{src}" alt=""></div>')
    return f'<img data-a="{anim}" class="s pic" style="--k:{k};{box_style(g)}" src="{src}" alt="">'


def unsupported(g, k, what):
    warnings.append(what)
    return f'<div data-a="box" class="s unsupported" style="--k:{k};{box_style(g)}">แปลงไม่ได้: {esc(what)}</div>'


def table_html(deck, shape, slide, g, k):
    tbl = q(shape._element, ".//a:tbl")
    cols = [deck.px(int(c.get("w"))) for c in tbl.findall("a:tblGrid/a:gridCol", NS)]
    tpr = q(tbl, "a:tblPr")
    styled = q(tpr, "a:tableStyleId") is not None             # สไตล์ตารางของ PowerPoint → ประมาณด้วย accent1: หัวเข้ม · แถวสลับอ่อน
    acc = "#" + deck.scheme.get("accent1", "4472C4")
    rows = []
    for ri, tr in enumerate(tbl.findall("a:tr", NS)):
        cells = []
        head = styled and ri == 0 and tpr.get("firstRow") == "1"
        band = styled and tpr.get("bandRow") == "1" and not head and (ri - (tpr.get("firstRow") == "1")) % 2 == 0
        for tc in tr.findall("a:tc", NS):
            if tc.get("hMerge") == "1" or tc.get("vMerge") == "1":
                continue
            pr = q(tc, "a:tcPr")
            css = []
            if pr is not None:
                css.append("padding:" + " ".join(f"{deck.px(int(pr.get(a, d)))}px" for a, d in
                                                  (("marT", 45720), ("marR", 91440), ("marB", 45720), ("marL", 91440))))
                css.append("vertical-align:" + {"ctr": "middle", "b": "bottom"}.get(pr.get("anchor"), "top"))
                fill = deck.color(q(pr, "a:solidFill"))
                if fill:
                    css.append(f"background:{fill}")
                for side, tag in (("left", "a:lnL"), ("right", "a:lnR"), ("top", "a:lnT"), ("bottom", "a:lnB")):
                    ln = q(pr, tag)
                    if ln is not None and q(ln, "a:noFill") is None and deck.color(ln):
                        css.append(f"border-{side}:{deck.px(int(ln.get('w', 12700))):.1f}px solid {deck.color(ln)}")
            if styled and not any("background" in c for c in css):
                css.append(f"background:{acc}" if head else f"background:color-mix(in srgb,{acc} {22 if band else 10}%,#fff)")
                css.append("border-bottom:2px solid #fff")
            span = "".join(f' {a}="{tc.get(b)}"' for a, b in (("colspan", "gridSpan"), ("rowspan", "rowSpan")) if tc.get(b))
            paras = []
            for para in tc.findall("a:txBody/a:p", NS):
                ppr = q(para, "a:pPr")
                algn = ppr.get("algn") if ppr is not None else None
                st = "text-align:" + {"ctr": "center", "r": "right"}[algn] if algn in ("ctr", "r") else ""
                spans = []
                for r in para.findall("a:r", NS):
                    rcss, _ = run_css(deck, shape, slide, para, q(r, "a:rPr"), 1, "#FFFFFF" if head else "#000000" if styled else None)
                    rcss += ";font-weight:700" if head else ""
                    spans.append(f'<span style="{rcss}">{esc(q(r, "a:t").text or "")}</span>')
                paras.append(f'<p style="{st}">{"".join(spans) or "&#8203;"}</p>')
            cells.append(f'<td{span} style="{";".join(css)}">{"".join(paras)}</td>')
        rows.append(f'<tr style="height:{deck.px(int(tr.get("h")))}px">{"".join(cells)}</tr>')
    colg = "".join(f'<col style="width:{w}px">' for w in cols)
    return f'<table data-a="table" class="s tbl" style="--k:{k};{box_style(g)}"><colgroup>{colg}</colgroup>{"".join(rows)}</table>'


def nice_axis(vmax: float) -> tuple[float, float]:
    """ค่าสูงสุดและระยะขีดที่เป็นเลขกลม (1 · 2 · 2.5 · 5 × 10^n) ประมาณ 5 ขีด"""
    raw = max(vmax, 1e-9) / 5
    mag = 10 ** math.floor(math.log10(raw))
    step = next(m * mag for m in (1, 2, 2.5, 5, 10) if m * mag >= raw)
    return math.ceil(vmax / step - 1e-9) * step, step


def chart_html(deck, shape, g, k):
    try:
        return _chart_html(deck, shape, g, k)
    except Exception as exc:                              # noqa: BLE001 — กราฟเดียวพังต้องไม่ทำให้ทั้งเด็คพัง
        return unsupported(g, k, f"กราฟ ({exc.__class__.__name__}: {exc})"[:80])


def _chart_html(deck, shape, g, k):
    try:
        ch = shape.chart
        plot = ch.plots[0]
        if len(ch.plots) > 1:
            warnings.append("กราฟ combo (แสดงแค่ชนิดแรก)")
        kind = str(ch.chart_type).split(".")[-1].split(" ")[0]
        cats = [str(c) for c in plot.categories]
        series = [(s.name, [v or 0 for v in s.values]) for s in plot.series]
        labels = plot.has_data_labels
        fmt = plot.data_labels.number_format if labels and not plot.data_labels.number_format_is_linked else ""
    except Exception as exc:                              # noqa: BLE001 — กราฟที่อ่านไม่ได้ไม่ควรทำให้ทั้งเด็คพัง
        try:
            kind = str(shape.chart.chart_type).split(" ")[0]
        except Exception:                                 # noqa: BLE001
            kind = "ไม่ทราบชนิด"
        return unsupported(g, k, f"กราฟชนิด {kind} (อ่านข้อมูลไม่ได้)")
    pal = [deck.color(q(s._element, "c:spPr")) or "#%s" % deck.scheme.get(f"accent{j + 1}", "4472C4")
           for j, s in enumerate(plot.series)]
    W, H = g["w"], g["h"]
    # ขนาดตัวอักษรของแต่ละส่วนตามที่ตั้งในกราฟ (pt × 2 = px) · ไม่ตั้ง = ขนาดทั้งกราฟ หรือ 16px
    cx = etree.fromstring(ch.part.blob)
    def fsz(part):
        el = cx.find(f".//c:{part}//a:defRPr[@sz]", NS)
        el = el if el is not None else cx.find("c:txPr//a:defRPr[@sz]", NS)
        return int(el.get("sz")) / 50 if el is not None else 16
    f_title, f_dl, f_cat, f_val, f_leg = fsz("title"), fsz("dLbls"), fsz("catAx"), fsz("valAx"), fsz("legend")
    title = ch.chart_title.text_frame.text if ch.has_title and ch.chart_title.has_text_frame else ""
    top, bottom, left, right = (48 if title else 12), 70, 56, 12
    vmax, step = nice_axis(max([max(v) for _, v in series if v] + [0]) or 1)
    try:                                                  # ค่าที่ตั้งไว้บนแกนใน PowerPoint ชนะค่าที่คำนวณ
        ax = ch.value_axis
        vmax = ax.maximum_scale or vmax
        step = ax.major_unit or (nice_axis(vmax)[1] if ax.maximum_scale else step)
    except (ValueError, AttributeError):
        pass
    dec = len(fmt.split(".")[1]) if "." in fmt else (2 if vmax <= 1 else 1 if any(x % 1 for _, v in series for x in v) else 0)
    tdec = 0 if step >= 1 else min(4, len(f"{step:.6f}".rstrip("0").split(".")[1]))
    num = lambda v, d=dec: f"{v:,.{d}f}"
    svg = [f'<svg class="chart" viewBox="0 0 {W} {H}" width="{W}" height="{H}" font-family="{FONT_STACK}">']
    if title:
        svg.append(f'<text x="{W / 2}" y="{f_title + 8:.0f}" text-anchor="middle" font-size="{f_title:.0f}" fill="#0E2A4D">{esc(title)}</text>')
    legend = ch.has_legend
    widths = [21 + 0.62 * f_leg * len(str(nm)) + 16 for nm, _ in series]
    x = (W - sum(widths)) / 2
    for j, (name, _) in enumerate(series if legend else []):
        svg.append(f'<rect x="{x:.1f}" y="{H - 26}" width="15" height="15" fill="{pal[j]}"/>'
                   f'<text x="{x + 21:.1f}" y="{H - 13}" font-size="{f_leg:.0f}" fill="#333">{esc(str(name))}</text>')
        x += widths[j]
    pw, ph = W - left - right, H - top - bottom
    if kind.startswith(("COLUMN", "BAR")):
        horiz = kind.startswith("BAR")
        n, m = len(cats), len(series)
        stacked = "STACKED" in kind
        if stacked:                                       # ซ้อนกัน: แกนตามผลรวมต่อหมวด · 100% = แปลงเป็นร้อยละก่อน
            if kind.endswith("100"):
                sums = [sum(v[ci] for _, v in series if ci < len(v)) or 1 for ci in range(n)]
                series = [(nm, [x / sums[ci] * 100 for ci, x in enumerate(v)]) for nm, v in series]
            vmax, step = nice_axis(max([sum(v[ci] for _, v in series if ci < len(v)) for ci in range(n)] + [0]))
            if kind.endswith("100"):
                vmax, step = 100, 20
        base = [0.0] * n
        try:                                              # กราฟแท่งนอน: PowerPoint วางหมวดแรกไว้ล่างสุด เว้นแต่กลับแกน
            flip = horiz and not ch.category_axis.reverse_order
        except (ValueError, AttributeError):
            flip = horiz
        row = (lambda ci: n - 1 - ci) if flip else (lambda ci: ci)
        band = (ph if horiz else pw) / max(1, n)
        bw = band * (0.6 if stacked else 0.7 / max(1, m))
        for t in range(round(vmax / step) + 1):
            v = step * t
            if horiz:
                x = left + pw * v / vmax
                svg.append(f'<line x1="{x:.1f}" x2="{x:.1f}" y1="{top}" y2="{top + ph}" stroke="#E5EAF0"/>'
                           f'<text x="{x:.1f}" y="{top + ph + 22}" text-anchor="middle" font-size="{f_val:.0f}" fill="#555">{num(v, tdec)}</text>')
            else:
                y = top + ph - ph * v / vmax
                svg.append(f'<line x1="{left}" x2="{W - right}" y1="{y:.1f}" y2="{y:.1f}" stroke="#E5EAF0"/>'
                           f'<text x="{left - 8}" y="{y + 6:.1f}" text-anchor="end" font-size="{f_val:.0f}" fill="#555">{num(v, tdec)}</text>')
        for ci, cat in enumerate(cats):
            for si, (_, vals) in enumerate(series):
                v = vals[ci] if ci < len(vals) else 0
                d = ci * m + si
                slot = 0.2 if stacked else 0.15 + si * bw / band
                if horiz:
                    y, w, x0 = top + row(ci) * band + band * slot, pw * v / vmax, left + pw * base[ci] / vmax
                    lx = x0 + w / 2 if stacked else x0 + w + 4
                    svg.append(f'<rect class="gb h" style="--d:{d}" x="{x0:.1f}" y="{y:.1f}" width="{w:.1f}" height="{bw:.1f}" fill="{pal[si]}"/>'
                               + (f'<text class="dl" style="--d:{d}" x="{lx:.1f}" y="{y + bw / 2 + 5:.1f}" text-anchor="{"middle" if stacked else "start"}" font-size="{f_dl:.0f}" fill="{"#fff" if stacked else "#333"}">{num(v)}</text>' if labels else ""))
                else:
                    x, h, y0 = left + ci * band + band * slot, ph * v / vmax, top + ph - ph * base[ci] / vmax
                    ly = y0 - h / 2 + 5 if stacked else y0 - h - 6
                    svg.append(f'<rect class="gb v" style="--d:{d}" x="{x:.1f}" y="{y0 - h:.1f}" width="{bw:.1f}" height="{h:.1f}" fill="{pal[si]}"/>'
                               + (f'<text class="dl" style="--d:{d}" x="{x + bw / 2:.1f}" y="{ly:.1f}" text-anchor="middle" font-size="{min(f_dl, max(11, bw * 0.45)):.0f}" fill="{"#fff" if stacked else "#333"}">{num(v)}</text>' if labels else ""))
                if stacked:
                    base[ci] += v
            if horiz:
                svg.append(f'<text x="{left - 6}" y="{top + row(ci) * band + band / 2 + 5:.1f}" text-anchor="end" font-size="{f_cat:.0f}" fill="#555">{esc(cat)}</text>')
            else:
                svg.append(f'<text x="{left + ci * band + band / 2:.1f}" y="{top + ph + 24}" text-anchor="middle" font-size="{f_cat:.0f}" fill="#333">{esc(cat)}</text>')
    elif kind.startswith("LINE"):
        for t in range(round(vmax / step) + 1):
            y = top + ph - ph * step * t / vmax
            svg.append(f'<line x1="{left}" x2="{W - right}" y1="{y:.1f}" y2="{y:.1f}" stroke="#E5EAF0"/>'
                       f'<text x="{left - 8}" y="{y + 6:.1f}" text-anchor="end" font-size="{f_val:.0f}" fill="#555">{num(step * t, tdec)}</text>')
        n = max(1, len(cats) - 1)
        for si, (_, vals) in enumerate(series):
            pts = " ".join(f"{left + 20 + (pw - 40) * ci / n:.1f},{top + ph - ph * v / vmax:.1f}" for ci, v in enumerate(vals))
            svg.append(f'<polyline class="dl" style="--d:{si}" points="{pts}" fill="none" stroke="{pal[si]}" stroke-width="4"/>')
        for ci, cat in enumerate(cats):
            svg.append(f'<text x="{left + 20 + (pw - 40) * ci / n:.1f}" y="{top + ph + 24}" text-anchor="middle" font-size="{f_cat:.0f}" fill="#333">{esc(cat)}</text>')
    elif kind.startswith(("PIE", "DOUGHNUT")):
        vals = series[0][1]
        pts = plot.series[0].points
        cols = [deck.color(q(pts[j].format._element, "c:spPr")) if j < len(pts) else None for j in range(len(vals))]
        cols = [c or "#%s" % deck.scheme.get(f"accent{j % 6 + 1}", "4472C4") for j, c in enumerate(cols)]
        svg = svg[:1 + bool(title)]                      # ทิ้ง legend ชื่อ series · pie ใช้ชื่อหมวดแทน
        cx, cy, r = W / 2, top + ph / 2 + 10, min(pw, ph) / 2
        total, a0 = sum(vals) or 1, -math.pi / 2
        for j, v in enumerate(vals):
            a1 = a0 + 2 * math.pi * v / total
            big = 1 if a1 - a0 > math.pi else 0
            x0, y0, x1, y1 = cx + r * math.cos(a0), cy + r * math.sin(a0), cx + r * math.cos(a1), cy + r * math.sin(a1)
            svg.append(f'<path class="dl" style="--d:{j}" d="M{cx:.1f},{cy:.1f} L{x0:.1f},{y0:.1f} A{r:.1f},{r:.1f} 0 {big} 1 {x1:.1f},{y1:.1f} Z" fill="{cols[j]}" stroke="#fff" stroke-width="2"/>')
            if labels:
                am = (a0 + a1) / 2
                svg.append(f'<text class="dl" style="--d:{j}" x="{cx + r * 0.65 * math.cos(am):.1f}" y="{cy + r * 0.65 * math.sin(am) + 6:.1f}" '
                           f'text-anchor="middle" font-size="{f_val:.0f}" fill="#fff">{num(v)}</text>')
            a0 = a1
        if kind.startswith("DOUGHNUT"):
            svg.append(f'<circle cx="{cx:.1f}" cy="{cy:.1f}" r="{r * 0.5:.1f}" fill="#fff"/>')
        x = left
        for j, cat in enumerate(cats if legend else []):
            svg.append(f'<rect x="{x:.1f}" y="{H - 26}" width="15" height="15" fill="{cols[j]}"/>'
                       f'<text x="{x + 21:.1f}" y="{H - 13}" font-size="{f_leg:.0f}" fill="#333">{esc(cat)}</text>')
            x += 21 + 0.62 * f_leg * len(cat) + 16
    else:
        return unsupported(g, k, f"กราฟชนิด {kind}")
    svg.append("</svg>")
    return f'<div data-a="chart" class="s chartbox" style="--k:{k};{box_style(g)}">{"".join(svg)}</div>'


def shape_html(deck, slide, part, shape, el, off, k, anim_title):
    name = etree.QName(el).localname
    g = geometry(deck, el, off, shape)
    if g is None:
        return ""
    if name == "cxnSp":
        return line_svg(deck, el, g, k)
    if name == "pic" and (q(el, ".//a:videoFile") is not None or q(el, ".//a:audioFile") is not None):
        return unsupported(g, k, "วิดีโอหรือเสียง")
    if name == "pic":
        return picture_html(deck, part, el, g, k)
    if name == "graphicFrame":
        uri = q(el, ".//a:graphicData").get("uri", "")
        if uri.endswith("/table"):
            return table_html(deck, shape, slide, g, k)
        if uri.endswith("/chart"):
            return chart_html(deck, shape, g, k)
        return unsupported(g, k, "SmartArt หรือวัตถุฝัง" if "diagram" in uri else "วัตถุฝัง")
    if name != "sp":
        return ""
    prst_el = q(el, "p:spPr/a:prstGeom")
    prst = prst_el.get("prst") if prst_el is not None else ("custom" if q(el, "p:spPr/a:custGeom") is not None else "rect")
    if prst in ("line", "straightConnector1") or prst.startswith("bentConnector"):
        return line_svg(deck, el, g, k)
    sppr = q(el, "p:spPr")
    if q(sppr, "a:blipFill") is not None or q(el, "p:blipFill") is not None:
        return picture_html(deck, part, el, g, k)
    out = []
    looks = fill_line(deck, el)
    if looks:
        if prst == "ellipse":
            looks.append("border-radius:50%")
        elif prst == "donut":
            adj = q(prst_el, "a:avLst/a:gd")
            ring = min(g["w"], g["h"]) * (int(re.sub(r"\D", "", adj.get("fmla"))) / 100000 if adj is not None else 0.25)
            looks = [f"border-radius:50%;border:{ring:.1f}px solid {paint(deck, el)[0] or '#808080'};background:transparent"]
        elif prst == "roundRect":
            adj = q(prst_el, "a:avLst/a:gd")
            ratio = int(re.sub(r"\D", "", adj.get("fmla"))) / 100000 if adj is not None else 0.16667
            looks.append(f"border-radius:{min(g['w'], g['h']) * ratio:.1f}px")
        if prst in POLY:
            fill, stroke, sw = paint(deck, el)
            fill = fill or "none"
            out.append(f'<div data-a="box" class="s b {prst}" style="--k:{k};{box_style(g)}"><svg class="geo" viewBox="0 0 100 100" preserveAspectRatio="none">'
                       f'<polygon points="{POLY[prst]}" fill="{fill}" stroke="{stroke or "none"}" stroke-width="{sw}" vector-effect="non-scaling-stroke"/></svg></div>')
        elif prst in ("rect", "roundRect", "ellipse", "donut", "snip1Rect", "round2SameRect", "flowChartProcess", "flowChartAlternateProcess"):
            out.append(f'<div data-a="box" class="s b {prst}" style="--k:{k};{box_style(g, ";" + ";".join(looks))}"></div>')
        else:
            warnings.append(f"รูปทรง {prst} (วาดเป็นสี่เหลี่ยม)")
            out.append(f'<div data-a="box" class="s b {prst}" style="--k:{k};{box_style(g, ";" + ";".join(looks))}"></div>')
    txt = text_html(deck, shape, slide)
    if txt and re.sub(r"<[^>]+>|&#8203;", "", txt).strip():
        a = anim_title if anim_title else ("pop" if looks and prst == "ellipse" else "up")
        kk = "" if anim_title == "title" else f"--k:{k + (0.35 if looks else 0)};"
        out.append(f'<div data-a="{a}" class="s t" style="{kk}{box_style(g)}">{txt}</div>')
    return "".join(out)


# ---------------------------------------------------------------- slide
def click_steps(slide) -> dict[str, int]:
    """spid → ขั้นที่ (1, 2, …) จาก animation แบบ On Click ใน mainSeq"""
    seq = slide._element.find(".//p:timing//p:cTn[@nodeType='mainSeq']", NS)
    steps = {}
    if seq is None:
        return steps
    for n, par in enumerate(seq.findall("p:childTnLst/p:par", NS), 1):
        for tgt in par.findall(".//p:spTgt", NS):
            steps.setdefault(tgt.get("spid"), n)
    return steps


def spid(el) -> str:
    c = el.find(".//p:cNvPr", NS)
    return c.get("id") if c is not None else ""


def flatten(shapes, off, groups=()):
    """ทุก shape ในสไลด์ รวมลูกในกลุ่ม · groups = id ของกลุ่มที่ครอบอยู่ (ขั้นของกลุ่มใช้กับลูกด้วย)"""
    for sh in shapes:
        if sh.shape_type is not None and str(sh.shape_type).startswith("GROUP"):
            yield from flatten(sh.shapes, Offset(off, sh._element), groups + (spid(sh._element),))
        else:
            yield sh, sh._element, off, groups


def stagger(items):
    """--k ตามลำดับการอ่าน: แถวละ 1 จังหวะ (บนลงล่าง) · ชิ้นในแถวเดียวกันเข้าพร้อมกัน"""
    rows, ks = [], []
    for g in items:
        if g is None:
            ks.append(0)
            continue
        for j, top in enumerate(rows):
            if abs(g["y"] - top) < 40:
                ks.append(j)
                break
        else:
            rows.append(g["y"])
            ks.append(len(rows) - 1)
    order = sorted(set(ks))
    return [order.index(x) for x in ks]


def sections(prs) -> dict[int, str]:
    lst = prs.part._element.find(".//p14:sectionLst", NS)
    ids = [s.slide_id for s in prs.slides]
    out = {}
    if lst is None:
        return out
    for sec in lst.findall("p14:section", NS):
        first = sec.find("p14:sldIdLst/p14:sldId", NS)
        if first is not None and int(first.get("id")) in ids:
            out[ids.index(int(first.get("id")))] = sec.get("name")
    return out


def background(deck, slide) -> str:
    for owner in (slide, slide.slide_layout, slide.slide_layout.slide_master):
        bg = q(owner._element, "p:cSld/p:bg")
        if bg is None:
            continue
        fill = deck.color(q(bg, "p:bgPr/a:solidFill"))
        if fill:
            return f"background:{fill}"
        ref = q(bg, "p:bgRef")
        if ref is not None and len(ref):
            return f"background:{deck.clr(ref[0])}"
        blip = q(bg, ".//a:blip")
        if blip is not None:
            img = owner.part.related_part(blip.get(f"{{{NS['r']}}}embed"))
            return f"background:url(data:{img.content_type};base64,{base64.b64encode(img.blob).decode()}) center/cover"
    return "background:#FFFFFF"


def slide_html(deck, slide, n, part_name):
    steps = click_steps(slide)
    items = list(flatten(slide.shapes, Offset()))
    geos = [geometry(deck, el, off, sh) for sh, el, off, _ in items]
    ks = stagger(geos)
    body = []
    title_text = ""
    for (sh, el, off, groups), k in zip(items, ks):
        anim_title = None
        if is_title(sh):
            title_text = title_text or sh.text_frame.text.strip()
            anim_title = "title" if "CENTER" not in str(sh.placeholder_format.type) else None
        piece = shape_html(deck, slide, slide.part, sh, el, off, k, anim_title)
        step = next((steps[x] for x in (spid(el),) + groups[::-1] if x in steps), None)
        if piece and step:
            piece = re.sub(r"<(div|svg|img|table) data-a=", rf'<\1 data-step="{step}" data-a=', piece)
        body.append(piece)
    if not title_text:
        texts = [s.text_frame.text.strip() for s, _, _, _ in items if s.has_text_frame and s.text_frame.text.strip()]
        title_text = texts[0] if texts else f"หน้า {n}"
    notes = ""
    if slide.has_notes_slide:
        t = slide.notes_slide.notes_text_frame.text.strip()
        if t:
            notes = "<aside>" + esc(t).replace("\n", "<br>") + "</aside>"
    part = f' data-part="{esc(part_name)}"' if part_name else ""
    menu = esc(title_text.splitlines()[0][:80])
    return f'<section class="slide" data-menu="{menu}"{part} style="{background(deck, slide)}">{"".join(body)}{notes}</section>'


CHROME = """<div id="menu"><div class="mhead"><b>สารบัญ<span>คลิกหน้าที่ต้องการ · กด M หรือ Esc เพื่อปิด</span></b><button class="mclose" type="button">✕ ปิด</button></div><div id="menuList"></div></div>
<div class="edge l" title="ก่อนหน้า (←)">‹</div><div class="edge r" title="ถัดไป (→)">›</div>
<div id="bar"></div><div id="bartip"></div>
<div id="notes"><button class="nclose" type="button" title="ปิดโน้ต (N)">✕</button><div id="notesBody"></div></div>
<div id="help"><b>← →</b> เปลี่ยนหน้า · <b>M</b> สารบัญ · <b>N</b> โน้ตผู้พูด · <b>T</b> จับเวลา · <b>F</b> เต็มจอ · <b>R</b> เล่นซ้ำ · <b>P</b> เล่นอัตโนมัติ</div>
<div id="ui"><span id="clock" title="เวลาที่ใช้ (กด T เริ่ม/หยุด)">0:00</span><button id="prev" title="ก่อนหน้า (←)">‹</button><span id="count"></span><button id="next" title="ถัดไป (→)">›</button><button id="mb" title="สารบัญ (M)">สารบัญ</button><button id="nb" title="โน้ตผู้พูด (N)">โน้ต</button><button id="fs" title="เต็มจอ (F)">เต็มจอ</button></div>"""


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("pptx", type=Path)
    ap.add_argument("-o", "--out", type=Path)
    ap.add_argument("--title", default="")
    ap.add_argument("--accent", default="", help="สีแถบความคืบหน้าและปุ่ม เช่น E4572E · ไม่ใส่ = accent1 ของธีม")
    ap.add_argument("--insert", action="append", default=[], metavar="N:file.html",
                    help='แทรกหน้าที่เขียนเอง (หน้าโต้ตอบ · demo) หลังหน้า N ของ .pptx · ไฟล์มี <section class="slide"> ได้หลายหน้า'
                         " และ <style> <script> ของตัวเอง · ใส่ซ้ำได้")
    a = ap.parse_args()
    deck = Deck(a.pptx)
    secs = sections(deck.prs)
    slides = [slide_html(deck, s, n, secs.get(n - 1, "" if n > 1 else "เริ่ม")) for n, s in enumerate(deck.prs.slides, 1)]
    extra_css, extra_js = [], []
    for spec in sorted(a.insert, key=lambda x: (-int(x.split(":", 1)[0]), -a.insert.index(x))):  # แทรกจากท้าย เลขหน้าเดิมไม่เลื่อน
        n, path = spec.split(":", 1)
        src = Path(path).read_text(encoding="utf-8")
        extra_css += re.findall(r"<style>(.*?)</style>", src, re.S)
        extra_js += re.findall(r"<script>(.*?)</script>", src, re.S)
        found = re.findall(r'<section class="slide.*?</section>', src, re.S)
        if not found:
            raise SystemExit(f'{path}: ไม่มี <section class="slide">')
        slides[int(n):int(n)] = [x.replace('<section class="slide', '<section data-inserted="1" class="slide', 1) for x in found]
    acc = "#" + (a.accent.lstrip("#") or deck.scheme.get("accent1", "E4572E"))
    markers = "".join(f'<marker id="m{c[1:]}" viewBox="0 0 10 10" refX="9" refY="5" markerWidth="5" markerHeight="5" orient="auto-start-reverse">'
                      f'<path d="M0,0 L10,5 L0,10 z" fill="{c}"/></marker>' for c in sorted(deck.markers))
    title = a.title or deck.prs.core_properties.title or a.pptx.stem
    css = (ASSETS / "runtime.css").read_text(encoding="utf-8")
    js = (ASSETS / "runtime.js").read_text(encoding="utf-8")
    out = (f'<!doctype html><html lang="th"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
           f'<title>{esc(title)}</title><style>{css}\n:root{{--acc:{acc};--acc2:{acc};--W:{deck.W}px;--H:{deck.H}px}}\n{"".join(extra_css)}</style></head><body>\n'
           f'<div id="stage"><svg width="0" height="0" style="position:absolute"><defs>{markers}</defs></svg>'
           + "\n".join(slides) + f"\n</div>\n{CHROME}\n<script>{js}</script>"
           + "".join(f"<script>{x}</script>" for x in extra_js) + "</body></html>")
    target = a.out or a.pptx.with_suffix(".html")
    target.write_text(out, encoding="utf-8")
    print(f"OK {target} · {len(slides)} หน้า · {len(out) // 1024} KB")
    for w in sorted(set(warnings)):
        print(f"  ⚠ แปลงไม่ครบ: {w} × {warnings.count(w)}")


if __name__ == "__main__":
    main()
