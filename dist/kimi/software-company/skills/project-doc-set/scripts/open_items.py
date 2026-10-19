"""open_items.py — รวมจุดที่ยังค้างในเอกสาร .docx ทั้งชุด (DRAFT) เป็นทะเบียน Excel ไฟล์เดียว

    python open_items.py docs/ docs/<CODE>-OPEN-ITEMS.xlsx --project <CODE>
    python open_items.py --self-test

หาในทุก .docx ของโฟลเดอร์:
  1. ตารางคำถาม      หัวตารางมี "คำถาม" (หรือ Question)            → 1 แถวต่อคำถามที่ยังไม่ปิด
  2. ตารางข้อสมมติ    หัวตารางมี "ผู้ยืนยัน" หรือ "สถานะ" + "ข้อสมมติ|ตัวเลข" → แถวที่ยังไม่ตัดสิน
  3. ตารางลงนาม      หัวตารางมี "บทบาท" + "สถานะ"                  → บทบาทที่ยัง "รอ"
  4. ป้ายในเนื้อหา    (รอยืนยัน) · (รอวัด…) · (รอ…) ที่อื่นทั้งหมด     → 1 แถวต่อย่อหน้าหรือแถวตาราง
ข้าม: ตารางประวัติการแก้ไข (หัวมี "เวอร์ชัน" + "แก้อะไร") · บรรทัดสารบัญ (style TOC) · บรรทัดที่อธิบายความหมายของป้าย ·
"(รอบหลัง)" · แถวที่ปิดแล้ว ("ตอบแล้ว" · "ปิดแล้ว" · "ตัดสินแล้ว" เว้นแต่ตามด้วย "ยังรอ" "ยังต้อง" ฯลฯ)
หัวข้อของแต่ละรายการอ่านจาก Heading style หรือ outline level ของย่อหน้า
ผล: แผ่น "สรุป" (นับด้วยสูตร COUNTIFS เปลี่ยนสถานะแล้วนับใหม่เอง) · แผ่น "รายการค้าง" (คอลัมน์ "สถานะ" "คำตอบ" สีเหลืองให้กรอก)
ต้องมี python-docx >= 1.1 และ openpyxl
"""
import argparse
import datetime
import pathlib
import re

from docx import Document
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

TAG = re.compile(r"รอยืนยัน|รอวัด|\(รอ(?!บ)[^)]{0,30}\)")   # "(รอบหลัง)" = later round, not open
CLOSED = re.compile(r"ตอบแล้ว|ปิดแล้ว|ตัดสินแล้ว|ไม่ต้องแก้|ไม่ใช่ขัดกัน|ไม่กระทบ")
STILL_OPEN = re.compile(r"ยังรอ|ยังต้อง|ต้องวัด|ยังไม่ตัดสิน|ยังไม่เลือก|ยังไม่มี|ยังไม่ตรวจ|รอเซ็น|รอเจ้าของ")
LEGEND = re.compile(r"ติดป้าย|ป้าย ?\(?รอ|วิธีอ่าน|\(รอยืนยัน\) · \(รอวัด|ฉบับร่างเพื่อทบทวน|^สรุป:|ผลตรวจชุด|ตารางนี้")
XREF = re.compile(r"\b(?:BRD |SRS |FSD )?(?:Q-[A-Z]*-?\d+|D\d{1,3}|AS-\d+)\b")
ORDER = ["PLAN", "BRD", "SRS", "FSD", "ARCH", "TESTPLAN"]


def text(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t"))).strip()


def rows_of(tbl):
    out = []
    for tr in tbl.iter(qn("w:tr")):
        cells = []
        for tc in tr.iter(qn("w:tc")):
            v = re.sub(r"\s+", " ", text(tc))
            if not cells or cells[-1] != v:      # merged cells repeat
                cells.append(v)
        out.append(cells)
    return out


def col(header, *names):
    for i, h in enumerate(header):
        if any(n in h for n in names):
            return i
    return None


def classify(owner, body):
    s = f"{owner} {body}"
    if re.search(r"รอวัด|ต้องวัด|วัดบน|วัดจริง", s):
        return "วัดจริง"
    if "กฎหมาย" in s or "DPO" in s:
        return "กฎหมาย"
    if re.search(r"เจ้าของผลิตภัณฑ์|เจ้าของ|JK|Owner", owner or "") or "เจ้าของผลิตภัณฑ์" in s:
        return "เจ้าของตัดสิน"
    if re.search(r"ตรวจ|review", s) and "โค้ด" in s:
        return "ตรวจโค้ด"
    return "ทีมตัดสิน"


def scan(path):
    code = re.sub(r"^.*?-([A-Z]+)-v.*$", r"\1", path.stem)
    items, heading = [], ""
    for el in Document(path).element.body.iterchildren():
        if el.tag == qn("w:p"):
            ppr = el.find(qn("w:pPr"))
            sty = ppr.find(qn("w:pStyle")).get(qn("w:val")) if ppr is not None and ppr.find(qn("w:pStyle")) is not None else ""
            t = text(el)
            lvl = ppr.find(qn("w:outlineLvl")) if ppr is not None else None
            if sty.lower().startswith("toc"):
                continue                                                # table of contents line
            if sty.lower().startswith("heading") or lvl is not None:
                heading = t
            elif TAG.search(t) and not LEGEND.search(t) and len(t) < 900:
                items.append(dict(doc=code, ref="", where=heading, issue=t, kind="", owner="", due="", interim="", impact=""))
            continue
        if el.tag != qn("w:tbl"):
            continue
        rows = rows_of(el)
        if not rows:
            continue
        hdr = rows[0]
        H = " ".join(hdr)
        if "เวอร์ชัน" in H and "แก้อะไร" in H:
            continue                                                    # revision history
        if "คำถาม" in H or "Question" in H:
            ci = col(hdr, "คำถาม", "Question")
            co = col(hdr, "ผู้ตอบ", "ส่งให้", "Owner")
            cd = col(hdr, "ตัดสินภายใน", "ภายใน", "Due")
            cf = col(hdr, "ตามอะไรไปก่อน", "ใช้ไปก่อน")
            cm = col(hdr, "ผลกระทบ", "กระทบ", "Impact")
            for r in rows[1:]:
                g = lambda i: r[i] if i is not None and i < len(r) else ""
                full = " ".join(r)
                if CLOSED.search(" ".join(r[:2])) or (CLOSED.search(full) and not STILL_OPEN.search(full)):
                    continue
                items.append(dict(doc=code, ref=r[0] if ci != 0 else "", where=heading, issue=g(ci), owner=g(co),
                                  due=g(cd), interim=g(cf), impact=g(cm), kind=""))
            continue
        if "ผู้ยืนยัน" in H or ("สถานะ" in H and re.search(r"ข้อสมมติ|ตัวเลข", H)):
            cs = col(hdr, "ผู้ยืนยัน", "สถานะ")
            for r in rows[1:]:
                st = r[cs] if cs is not None and cs < len(r) else ""
                if CLOSED.search(st) and not STILL_OPEN.search(st):
                    continue
                items.append(dict(doc=code, ref=r[0], where=heading, issue=r[1] if len(r) > 1 else "", owner=st,
                                  due="", interim="", impact=r[2] if len(r) > 2 else "", kind=""))
            continue
        if "บทบาท" in H and "สถานะ" in H:
            cs = col(hdr, "สถานะ")
            for r in rows[1:]:
                if cs is not None and cs < len(r) and "รอ" in r[cs]:
                    items.append(dict(doc=code, ref="ลงนาม", where=heading, issue=f"ลงนามรับเอกสาร: {r[0]}",
                                      owner=r[1] if len(r) > 1 else "", due="", interim="", impact="", kind="ลงนาม"))
            continue
        for r in rows[1:]:                                              # any other table: inline tags
            full = " | ".join(r)
            if TAG.search(full) and not LEGEND.search(full):
                items.append(dict(doc=code, ref=r[0][:40], where=heading, issue=full, owner="", due="",
                                  interim="", impact="", kind=""))
    for it in items:
        it["kind"] = it["kind"] or classify(it["owner"], it["issue"] + " " + it["due"])
        it["xref"] = " · ".join(sorted(set(m.strip() for m in XREF.findall(it["issue"]) if m.strip() != it["ref"])))
    return items


def link(items):
    """an item that points at another document's question takes that question's kind (owner decision wins)"""
    by_ref = {f'{it["doc"]} {it["ref"]}': it for it in items if it["ref"]}
    by_ref.update({it["ref"]: it for it in items if it["ref"]})
    for it in items:
        for x in filter(None, it["xref"].split(" · ")):
            tgt = by_ref.get(x) or by_ref.get(x.split(" ")[-1])
            if tgt and tgt is not it and tgt["kind"] == "เจ้าของตัดสิน":
                it["kind"] = "เจ้าของตัดสิน"
    return items


def write(items, out, project):
    wb = Workbook()
    F = "Arial"
    thin = Side(style="thin", color="D0D0D0")
    head_fill, input_fill = PatternFill("solid", fgColor="1F1F1F"), PatternFill("solid", fgColor="FFF7CC")
    kinds = ["เจ้าของตัดสิน", "ทีมตัดสิน", "กฎหมาย", "วัดจริง", "ตรวจโค้ด", "ลงนาม"]

    ws = wb.active
    ws.title = "รายการค้าง"
    cols = [("#", 5), ("เอกสาร", 10), ("รหัส", 12), ("หัวข้อในเอกสาร", 24), ("เรื่องที่ค้าง", 70), ("ประเภท", 14),
            ("ใครตอบ", 20), ("ตัดสินภายใน", 18), ("ใช้ไปก่อน / ค่าที่เสนอ", 30), ("กระทบ", 26), ("อ้างถึง", 14),
            ("สถานะ", 10), ("คำตอบ (กรอก)", 36)]
    for i, (h, w) in enumerate(cols, 1):
        c = ws.cell(1, i, h)
        c.font, c.fill = Font(name=F, bold=True, color="FFFFFF", size=10), head_fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    items.sort(key=lambda x: (ORDER.index(x["doc"]) if x["doc"] in ORDER else 99, kinds.index(x["kind"])))
    for n, it in enumerate(items, 1):
        vals = [n, it["doc"], it["ref"], it["where"], it["issue"], it["kind"], it["owner"], it["due"], it["interim"],
                it["impact"], it["xref"], "เปิด", ""]
        for i, v in enumerate(vals, 1):
            c = ws.cell(n + 1, i, v)
            c.font = Font(name=F, size=10)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        ws.cell(n + 1, 13).fill = input_fill
        ws.cell(n + 1, 12).fill = input_fill
    last = len(items) + 1
    dv = DataValidation(type="list", formula1='"เปิด,ตอบแล้ว,เลื่อน,ยกเลิก"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"L2:L{last}")
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:M{last}"

    s = wb.create_sheet("สรุป", 0)
    s.column_dimensions["A"].width = 16
    for i in range(2, 10):
        s.column_dimensions[get_column_letter(i)].width = 13
    s["A1"] = f"{project} — จุดที่ยังค้างในเอกสาร DRAFT"
    s["A1"].font = Font(name=F, bold=True, size=14)
    s["A2"] = f"สแกนเมื่อ {datetime.date.today():%d/%m/%Y} จากไฟล์ .docx ทุกฉบับ · นับเฉพาะแถวสถานะ \"เปิด\" (แก้สถานะได้ในแผ่น รายการค้าง)"
    s["A2"].font = Font(name=F, size=9, color="6B6B6B")
    docs = [d for d in ORDER if any(it["doc"] == d for it in items)] + sorted({it["doc"] for it in items} - set(ORDER))
    hdr = ["เอกสาร"] + kinds + ["รวม"]
    for j, h in enumerate(hdr, 1):
        c = s.cell(4, j, h)
        c.font, c.fill = Font(name=F, bold=True, color="FFFFFF", size=10), head_fill
        c.alignment = Alignment(horizontal="center")
    R = "'รายการค้าง'!"
    for i, d in enumerate(docs + ["รวม"], 5):
        s.cell(i, 1, d).font = Font(name=F, bold=(d == "รวม"), size=10)
        for j, k in enumerate(kinds, 2):
            L = get_column_letter(j)
            f = (f'=SUM({L}5:{L}{i - 1})' if d == "รวม" else
                 f'=COUNTIFS({R}$B$2:$B${last},$A{i},{R}$F$2:$F${last},{L}$4,{R}$L$2:$L${last},"เปิด")')
            s.cell(i, j, f).font = Font(name=F, size=10, bold=(d == "รวม"))
        T = get_column_letter(len(kinds) + 2)
        s.cell(i, len(kinds) + 2, f"=SUM(B{i}:{get_column_letter(len(kinds) + 1)}{i})").font = Font(name=F, bold=True, size=10)
        for j in range(1, len(hdr) + 1):
            s.cell(i, j).border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if j > 1:
                s.cell(i, j).alignment = Alignment(horizontal="center")
    r0 = 5 + len(docs) + 2
    notes = [("ประเภท", ""),
             ("เจ้าของตัดสิน", "เจ้าของผลิตภัณฑ์ต้องเลือกเอง ค้นจากเว็บแล้วปิดไม่ได้"),
             ("ทีมตัดสิน", "BA · SA · ผู้ทำ mockup ตัดสินได้ในทีม"),
             ("กฎหมาย", "รอฝ่ายกฎหมายหรือ DPO"),
             ("วัดจริง", "ต้องวัดบนระบบจริง (ส่วนใหญ่ใน Phase 0) ไม่มีตัวเลขเผยแพร่"),
             ("ตรวจโค้ด", "ต้องเปิดโค้ดตรวจ"),
             ("ลงนาม", "บทบาทที่ยังไม่ได้เซ็นรับเอกสาร"),
             ("วิธีใช้", "กรอกคอลัมน์ \"คำตอบ\" แล้วเปลี่ยน \"สถานะ\" เป็น ตอบแล้ว (ช่องสีเหลือง) · ตัวเลขในตารางนี้คำนวณใหม่เอง")]
    for k, (a, b) in enumerate(notes):
        s.cell(r0 + k, 1, a).font = Font(name=F, bold=True, size=10)
        s.cell(r0 + k, 2, b).font = Font(name=F, size=10)
    wb.save(out)
    return len(items)


def self_test():
    import tempfile
    from docx import Document as New
    tmp = pathlib.Path(tempfile.mkdtemp())
    d = New()
    d.add_heading("1. ขอบเขต", 1)
    d.add_paragraph("เวลาติดตั้ง 30 นาที (รอวัดใน Phase 0)")
    d.add_paragraph("ฟีเจอร์ปลุกเครื่อง (รอบหลัง)")                       # not open
    d.add_paragraph("วิธีอ่าน: ตัวเลขที่ยังไม่แน่ใจติดป้าย (รอยืนยัน)")       # legend
    t = d.add_table(rows=4, cols=3)
    for r, vals in enumerate([("รหัส", "คำถาม", "ผู้ตอบ"), ("Q-01", "คิดราคาต่อผู้ใช้หรือต่อเครื่อง", "เจ้าของผลิตภัณฑ์"),
                              ("Q-02", "ตอบแล้ว: ใช้ BSD-3", "เจ้าของผลิตภัณฑ์"), ("Q-03", "relay ตั้งที่ไหน (BRD Q-01)", "SA")]):
        for c, v in enumerate(vals):
            t.cell(r, c).text = v
    h = d.add_table(rows=2, cols=3)
    for c, v in enumerate(("เวอร์ชัน", "แก้อะไร", "ผู้แก้")):
        h.cell(0, c).text = v
    h.cell(1, 1).text = "ปิดป้าย (รอยืนยัน) 5 ข้อ"                         # revision history
    d.save(tmp / "X-SRS-v0.1-DRAFT.docx")
    items = link(scan(tmp / "X-SRS-v0.1-DRAFT.docx"))
    got = sorted((it["ref"] or it["issue"].split()[0], it["kind"]) for it in items)
    want = sorted([("Q-01", "เจ้าของตัดสิน"), ("Q-03", "เจ้าของตัดสิน"), ("เวลาติดตั้ง", "วัดจริง")])
    assert got == want, got
    assert items[0]["where"] == "1. ขอบเขต", items[0]["where"]
    assert write(items, tmp / "o.xlsx", "X") == 3
    print("self-test ok")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("docs", nargs="?")
    ap.add_argument("out", nargs="?")
    ap.add_argument("--project", default="")
    ap.add_argument("--self-test", action="store_true")
    a = ap.parse_args()
    if a.self_test:
        self_test()
        raise SystemExit
    if not (a.docs and a.out):
        ap.error("ต้องระบุ <docs_dir> <out.xlsx>")
    files = sorted(p for p in pathlib.Path(a.docs).glob("*.docx") if not p.name.startswith("~$"))
    allitems = link([it for f in files for it in scan(f)])
    n = write(allitems, a.out, a.project or files[0].stem.split("-")[0])
    print(f"{n} open items from {len(files)} files -> {a.out}")
