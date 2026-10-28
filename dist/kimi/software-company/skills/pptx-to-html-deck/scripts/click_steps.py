"""click_steps.py — ใส่ animation แบบ On Click (Fade) ให้ shape ในสไลด์ ที่สร้างด้วย python-pptx

    from click_steps import add_click_steps
    add_click_steps(slide, [[box1, label1], [box2], [arrow, box3]])   # 3 ขั้น · ขั้นละกี่ชิ้นก็ได้

PowerPoint เล่นเป็นคลิกทีละขั้น และ pptx_to_html.py แปลงเป็นขั้นในหน้า (→ เปิดทีละขั้น)
เรียกซ้ำกับสไลด์เดิม = เขียนทับลำดับเดิม

    python click_steps.py --self-test      # สร้าง .pptx 3 หน้า (มี 2 ขั้นในหน้า 2) แล้วอ่านขั้นกลับมานับ
ต้องมี python-pptx (ไม่มี → self-test ข้ามและบอกชัด)
"""
from __future__ import annotations

from lxml import etree

P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def _el(tag: str, parent=None, **attrs):
    e = etree.SubElement(parent, f"{{{P}}}{tag}") if parent is not None else etree.Element(f"{{{P}}}{tag}")
    for k, v in attrs.items():
        e.set(k, str(v))
    return e


def add_click_steps(slide, steps: list[list]) -> None:
    """steps = รายการขั้น · แต่ละขั้นคือรายการ shape ที่โผล่พร้อมกันเมื่อคลิก"""
    sld = slide._element
    old = sld.find(f"{{{P}}}timing")
    if old is not None:
        sld.remove(old)
    ids = iter(range(1, 10_000))
    timing = _el("timing")
    root = _el("cTn", _el("par", _el("tnLst", timing)), id=next(ids), dur="indefinite", restart="never", nodeType="tmRoot")
    seq = _el("seq", _el("childTnLst", root), concurrent=1, nextAc="seek")
    main = _el("childTnLst", _el("cTn", seq, id=next(ids), dur="indefinite", nodeType="mainSeq"))
    for shapes in steps:
        click = _el("cTn", _el("par", main), id=next(ids), fill="hold")
        _el("cond", _el("stCondLst", click), delay="indefinite")
        inner = _el("cTn", _el("par", _el("childTnLst", click)), id=next(ids), fill="hold")
        _el("cond", _el("stCondLst", inner), delay=0)
        group = _el("childTnLst", inner)
        for n, shape in enumerate(shapes):
            spid = str(shape.shape_id)
            eff = _el("cTn", _el("par", group), id=next(ids), presetID=10, presetClass="entr", presetSubtype=0,
                      fill="hold", nodeType="clickEffect" if n == 0 else "withEffect")
            _el("cond", _el("stCondLst", eff), delay=0)
            kids = _el("childTnLst", eff)
            beh = _el("cBhvr", _el("set", kids))
            _el("cond", _el("stCondLst", _el("cTn", beh, id=next(ids), dur=1, fill="hold")), delay=0)
            _el("spTgt", _el("tgtEl", beh), spid=spid)
            _el("attrName", _el("attrNameLst", beh)).text = "style.visibility"
            _el("strVal", _el("to", beh.getparent()), val="visible")
            fade = _el("cBhvr", _el("animEffect", kids, transition="in", filter="fade"))
            _el("cTn", fade, id=next(ids), dur=500)
            _el("spTgt", _el("tgtEl", fade), spid=spid)
    for evt, tag in (("onPrev", "prevCondLst"), ("onNext", "nextCondLst")):
        _el("sldTgt", _el("tgtEl", _el("cond", _el(tag, seq), evt=evt, delay=0)))
    # ลำดับใน p:sld: cSld · clrMapOvr · transition (อาจห่อใน mc:AlternateContent) · timing · extLst
    ext = sld.find(f"{{{P}}}extLst")
    if ext is not None:
        ext.addprevious(timing)
    else:
        sld.append(timing)


def sample_deck(path):
    """.pptx ตัวอย่างสำหรับ --self-test ของสคริปต์ในโฟลเดอร์นี้ · หน้า 2 มี 2 ขั้น (ขั้น 1 = 2 ชิ้น, ขั้น 2 = 1 ชิ้น)"""
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    for n in (1, 2, 3):
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = f"Slide {n}"
        if n == 2:
            boxes = [s.shapes.add_textbox(Inches(1), Inches(1.5 + k), Inches(4), Inches(0.6)) for k in range(3)]
            for k, b in enumerate(boxes):
                b.text_frame.text = f"Step box {k}"
            add_click_steps(s, [[boxes[0], boxes[1]], [boxes[2]]])
    prs.save(str(path))
    return path


def self_test() -> int:
    import tempfile
    try:
        from pptx import Presentation
    except ImportError:
        print("SKIP self-test: ไม่มี python-pptx (pip install python-pptx)")
        return 0
    from pptx_to_html import click_steps
    with tempfile.TemporaryDirectory() as tmp:
        prs = Presentation(str(sample_deck(f"{tmp}/t.pptx")))
    got = [click_steps(s) for s in prs.slides]
    checks = {
        "3 หน้า": len(got) == 3,
        "หน้า 1 และ 3 ไม่มีขั้น": not got[0] and not got[2],
        "หน้า 2 มี 3 ชิ้น": len(got[1]) == 3,
        "หน้า 2 มี 2 ขั้น (1 กับ 2)": sorted(set(got[1].values())) == [1, 2],
        "ขั้น 1 มี 2 ชิ้น ขั้น 2 มี 1 ชิ้น": sorted(list(got[1].values()).count(k) for k in (1, 2)) == [1, 2],
    }
    for name, ok in checks.items():
        print("ok  " if ok else "FAIL", name)
    print(f"self-test {sum(checks.values())}/{len(checks)}")
    return 0 if all(checks.values()) else 1


if __name__ == "__main__":
    import sys
    sys.exit(self_test() if "--self-test" in sys.argv else print(__doc__))
