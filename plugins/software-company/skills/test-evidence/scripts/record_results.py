"""Write test results into a test-case workbook someone already has, adding Result · Test result · Evidence columns.

  python record_results.py cases.xlsx results.csv --evidence-dir qa/evidence/2026-10-18-r1 -o qa/PRJ-test-cases.xlsx
  python record_results.py --self-test

- results.csv comes from make_evidence.py (columns id, result, test_result, evidence).
- Every sheet is searched. The header row is the first row in the top 10 with an ID column
  (id · tc id · test case id · case id · รหัส · เลขที่เคส). Rows are matched on that ID.
- A column that already exists (Result/Status/ผล, Test result/Actual/ผลที่ได้, Evidence/หลักฐาน) is reused.
  Missing ones are added at the end of the header, so formulas that point at existing columns keep working.
- Evidence becomes a clickable link, relative to the output workbook so the folder can be moved together.
- Result gets a Pass/Fail/Blocked/Not run dropdown and colour. IDs in results.csv not found in the workbook are listed.
Needs: pip install openpyxl
"""
import argparse, copy, csv, os, sys
from pathlib import Path
import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

RESULTS = {'Pass': 'D1FAE5', 'Fail': 'FEE2E2', 'Blocked': 'FEF3C7', 'Not run': 'EEEEF2'}
ID_NAMES = {'id', 'tc id', 'tc_id', 'test case id', 'testcase id', 'case id', 'test id', 'รหัส', 'เลขที่เคส', 'รหัสเคส'}
COLS = {  # key → (header to add, names that already mean the same thing)
    'result': ('Result', {'result', 'status', 'test status', 'ผล', 'ผลทดสอบ', 'สถานะ'}),
    'test_result': ('Test result', {'test result', 'test_result', 'actual', 'actual result', 'ผลที่ได้', 'ผลที่ได้จริง'}),
    'evidence': ('Evidence', {'evidence', 'หลักฐาน', 'screenshot', 'ภาพหลักฐาน'}),
}
WIDTH = {'result': 11, 'test_result': 45, 'evidence': 24}


def norm(v):
    return ' '.join(str(v or '').replace('\n', ' ').split()).lower()


def find_header(ws):
    for r in range(1, min(ws.max_row, 10) + 1):
        names = {norm(c.value): c.column for c in ws[r] if c.value not in (None, '')}
        for n, col in names.items():
            if n in ID_NAMES:
                return r, col, names
    return None


def record(book, results_csv, evidence_dir, out):
    rows = {r['id'].strip(): r for r in csv.DictReader(open(results_csv, encoding='utf-8-sig'))}
    wb = openpyxl.load_workbook(book)
    found, sheets = set(), 0
    for ws in wb.worksheets:
        hit = find_header(ws)
        if not hit:
            continue
        sheets += 1
        hr, id_col, names = hit
        last = max(c.column for c in ws[hr] if c.value not in (None, ''))
        style_src = ws.cell(hr, id_col)
        col = {}
        for key, (label, aliases) in COLS.items():
            col[key] = next((c for n, c in names.items() if n in aliases), None)
            if col[key] is None:
                last += 1
                cell = ws.cell(hr, last, label)
                cell.font, cell.fill, cell.border = copy.copy(style_src.font), copy.copy(style_src.fill), copy.copy(style_src.border)
                cell.alignment = copy.copy(style_src.alignment)
                ws.column_dimensions[get_column_letter(last)].width = WIDTH[key]
                col[key] = last
        end = hr
        for r in range(hr + 1, ws.max_row + 1):
            case = str(ws.cell(r, id_col).value or '').strip()
            if case:
                end = r
            if case not in rows:
                continue
            found.add(case)
            res = rows[case]
            ws.cell(r, col['result'], res.get('result', ''))
            t = ws.cell(r, col['test_result'], res.get('test_result', ''))
            t.alignment = Alignment(vertical='top', wrap_text=True)
            ev = res.get('evidence', '')
            cell = ws.cell(r, col['evidence'], ev)
            if ev:
                target = Path(evidence_dir) / ev if evidence_dir else Path(ev)
                cell.hyperlink = Path(os.path.relpath(target.resolve(), Path(out).resolve().parent)).as_posix()
                cell.font = Font(color='2563EB', underline='single')
        if end > hr:
            L = get_column_letter(col['result'])
            rng = f'{L}{hr + 1}:{L}{end}'
            dv = DataValidation(type='list', formula1='"' + ','.join(RESULTS) + '"', allow_blank=True)
            dv.add(rng); ws.add_data_validation(dv)
            for word, color in RESULTS.items():
                ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${L}{hr + 1}="{word}"'],
                                              fill=PatternFill('solid', fgColor=color, bgColor=color), font=Font(bold=True)))
    if not sheets:
        sys.exit('no sheet has an ID column in its first 10 rows (id · tc id · test case id · รหัส)')
    wb.save(out)
    missing = sorted(set(rows) - found)
    print(f'{out}: {len(found)} of {len(rows)} results written in {sheets} sheet(s)')
    if missing:
        print('not found in workbook: ' + ', '.join(missing))


def self_test():
    import tempfile
    tmp = Path(tempfile.mkdtemp())
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Login'
    ws.append(['Login tests']); ws.append([])
    ws.append(['TC ID', 'Title', 'Status']); ws.append(['TC-LOGIN-001', 'ok', None]); ws.append(['TC-LOGIN-003', 'bad', None])
    wb.save(tmp / 'in.xlsx')
    (tmp / 'ev').mkdir()
    with open(tmp / 'ev' / 'results.csv', 'w', encoding='utf-8-sig', newline='') as f:
        csv.writer(f).writerows([['id', 'result', 'test_result', 'evidence'], ['TC-LOGIN-001', 'Pass', 'เข้าได้', 'TC-LOGIN-001.gif'],
                                 ['TC-LOGIN-003', 'Fail', 'บอกว่าอีเมลมีจริง', 'TC-LOGIN-003.gif'], ['TC-X-9', 'Pass', '', '']])
    record(tmp / 'in.xlsx', tmp / 'ev' / 'results.csv', tmp / 'ev', tmp / 'out.xlsx')
    ws = openpyxl.load_workbook(tmp / 'out.xlsx').active
    assert [c.value for c in ws[3]] == ['TC ID', 'Title', 'Status', 'Test result', 'Evidence'], [c.value for c in ws[3]]
    assert ws['C5'].value == 'Fail' and ws['D5'].value == 'บอกว่าอีเมลมีจริง'
    assert ws['E4'].hyperlink.target == 'ev/TC-LOGIN-001.gif', ws['E4'].hyperlink.target
    print('self-test passed')


if __name__ == '__main__':
    if '--self-test' in sys.argv:
        self_test(); sys.exit()
    ap = argparse.ArgumentParser()
    ap.add_argument('workbook'); ap.add_argument('results_csv')
    ap.add_argument('--evidence-dir', help='folder the evidence file names in results.csv sit in (default: next to results.csv)')
    ap.add_argument('-o', '--out', required=True)
    a = ap.parse_args()
    record(a.workbook, a.results_csv, a.evidence_dir or Path(a.results_csv).parent, a.out)
