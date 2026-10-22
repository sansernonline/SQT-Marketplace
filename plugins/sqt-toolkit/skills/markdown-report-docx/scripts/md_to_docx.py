"""md_to_docx.py — เปลี่ยนรายงาน Markdown 1 ไฟล์เป็น .docx ที่จัดหน้าแล้ว (ปกโลโก้กลาง · สารบัญ · ตารางสีแบรนด์)

    python md_to_docx.py report.md                     # → report.docx ข้างไฟล์ต้นทาง
    python md_to_docx.py report.md -o _to_delete/Report_v1.0.docx --pdf

Markdown คือต้นฉบับ แก้ที่ .md แล้วรันใหม่ · ไม่แก้ .docx ด้วยมือ
ต้องมี python-docx >= 1.1 · ใช้ brandkit.py จาก skill branded-document-design (โฟลเดอร์ข้างกัน)
รูปแบบ front matter และสิ่งที่แต่ละบรรทัดกลายเป็น → SKILL.md ของ skill นี้
"""
from __future__ import annotations

import argparse
import os
import re
import subprocess
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parents[1] / "branded-document-design" / "scripts"))
from brandkit import SIZE, BrandDoc, keep_with_next, style_run, use_brand  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402

COLOR_KEYS = ("brand", "brand_2", "brand_deep", "brand_tint", "brand_tint_2")
CALLOUTS = {"NOTE": ("note", "หมายเหตุ"), "TIP": ("tip", "ข้อแนะนำ"), "IMPORTANT": ("warning", "สำคัญ"),
            "WARNING": ("warning", "ข้อควรระวัง"), "CAUTION": ("critical", "อันตราย")}
TABLE_ROW = re.compile(r"^\|.*\|\s*$")
TABLE_RULE = re.compile(r"^\|[-| :]+\|\s*$")


# -- อ่านต้นฉบับ -----------------------------------------------------------------------------------
def split_front(md: str) -> tuple[dict, str]:
    """front matter แบบ key: value บรรทัดละค่า → (meta, เนื้อหา)"""
    m = re.match(r"---\r?\n(.*?)\r?\n---\r?\n(.*)", md, re.S)
    if not m:
        return {}, md
    meta = dict(line.split(": ", 1) for line in m.group(1).splitlines() if ": " in line)
    return {k.strip(): v.strip() for k, v in meta.items()}, m.group(2)


def inline(text: str) -> list[tuple[str, dict]]:
    """**ตัวหนา** และ `โค้ด` → ชิ้นส่วนสำหรับ BrandDoc.rich()"""
    parts = []
    for piece in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if piece.startswith("**"):
            parts.append((piece[2:-2], {"bold": True, "color": "text"}))
        elif piece.startswith("`"):
            parts.append((piece[1:-1], {"mono": True, "color": "brand_deep"}))
        elif piece:
            parts.append((piece, {"color": "text_body"}))
    return parts


def plain(text: str) -> str:
    return re.sub(r"\*\*|`", "", text)


def widths(rows: list[list[str]], total=9360) -> list[int]:
    """กว้างตามข้อความยาวสุดของแต่ละคอลัมน์ แบบบีบช่วง → คอลัมน์รหัสและชื่อไม่ถูกตัดกลางคำ"""
    longest = [max(min(len(r[i]) if i < len(r) else 0, 120) for r in rows) for i in range(len(rows[0]))]
    raw = [max(n, 14) ** 0.7 for n in longest]
    raw[0] *= 1.35                                   # คอลัมน์แรกเป็นชื่อหรือหัวข้อ ชื่อไทยต้องไม่ขึ้นบรรทัดใหม่
    w = [int(total * n / sum(raw)) for n in raw]
    w[-1] += total - sum(w)
    return w


def next_line(lines: list[str], i: int) -> str:
    return next((x for x in lines[i + 1:] if x.strip()), "")


# -- ส่วนประกอบเฉพาะรายงาน ------------------------------------------------------------------------
def centered_cover(doc: BrandDoc, meta: dict, base: Path) -> None:
    """ปกแบบรายงาน: โลโก้กลางหน้า → ชื่อสีแบรนด์ → ชื่อรอง → บรรทัด meta → หมายเหตุตัวเอียง"""
    d = doc.doc
    d.add_paragraph().paragraph_format.space_after = Pt(float(meta.get("cover_top_pt", 150)))

    def line(text, size, color, bold=False, italic=False, after=4):
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(after)
        style_run(p.add_run(text), size=size, color=color, bold=bold, italic=italic)

    if meta.get("logo"):
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(14)
        p.add_run().add_picture(str(base / meta["logo"]), width=Cm(float(meta.get("logo_width_cm", 6.5))))
    line(meta.get("title", ""), SIZE["cover_title"], "brand", bold=True)
    if meta.get("subtitle"):
        line(meta["subtitle"], SIZE["cover_subtitle"], "text", bold=True, after=18)
    if meta.get("meta"):
        line(meta["meta"], SIZE["cover_meta"], "text_body", after=2)
    if meta.get("note"):
        line(meta["note"], SIZE["cover_note"], "text_muted", italic=True)
    d.add_page_break()


def numbered(doc: BrandDoc, items: list[str]) -> None:
    """ใส่เลขเอง เพราะ List Number ของ Word นับต่อข้ามรายการ"""
    for k, item in enumerate(items, 1):
        p = doc.rich([(f"{k}.", {"bold": True, "color": "brand"}), ("\t", {})] + inline(item), space_after=2)
        p.paragraph_format.left_indent = Cm(0.9)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))


def bullets(doc: BrandDoc, items: list[str]) -> None:
    for item in items:
        p = doc.doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        for text, opts in inline(item):
            style_run(p.add_run(text), size=SIZE["body"], **opts)


def keep_table_together(table, rows: int) -> None:
    """ตารางสั้น (≤ 8 แถว) อยู่หน้าเดียวกันทั้งตาราง"""
    if rows <= 8:
        for row in table.rows[:-1]:
            for cell in row.cells:
                for par in cell.paragraphs:
                    keep_with_next(par)


# -- แปลงทีละบรรทัด -------------------------------------------------------------------------------
def build(md_path: Path, out: Path) -> Path:
    meta, body = split_front(md_path.read_text(encoding="utf-8"))
    use_brand(**{k: meta[k] for k in COLOR_KEYS if k in meta})
    doc = BrandDoc(footer_text=meta.get("footer", "หน้า"))
    base = md_path.parent
    centered_cover(doc, meta, base)
    if meta.get("toc", "yes") != "no":
        doc.toc()

    lines, i, fig = body.splitlines(), 0, 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line or line == "---" or (line.startswith("<!--") and "pagebreak" not in line):
            i += 1
            continue
        if "<!-- pagebreak -->" in line:
            doc.page_break()
        elif line.startswith("```"):
            code, i = [], i + 1
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            doc.code("\n".join(code))
        elif m := re.match(r"(#{1,3}) (.+)", line):
            getattr(doc, f"h{len(m.group(1))}")(plain(m.group(2)))
        elif line.startswith(">"):
            quote = []
            while i < len(lines) and lines[i].startswith(">"):
                quote.append(lines[i].lstrip("> ").rstrip())
                i += 1
            kind, title = "note", "หมายเหตุ"
            if m := re.match(r"\[!(\w+)\]\s*(.*)", quote[0]):
                kind, title = CALLOUTS.get(m.group(1).upper(), CALLOUTS["NOTE"])
                quote[0] = m.group(2)
            t = doc.callout(kind, title, plain(" ".join(q for q in quote if q)))
            if TABLE_ROW.match(next_line(lines, i - 1)):
                for par in t.rows[0].cells[0].paragraphs:
                    keep_with_next(par)
            continue
        elif m := re.match(r"!\[(.*?)\]\((.*?)\)", line):
            fig += 1
            cap = doc.figure(str(base / m.group(2)), m.group(1) or None,
                             width_cm=float(meta.get("figure_width_cm", 17)), number=fig)
            keep_with_next(cap)                              # คำบรรยายอยู่หน้าเดียวกับรูป
        elif TABLE_ROW.match(line):
            block = []
            while i < len(lines) and TABLE_ROW.match(lines[i]):
                if not TABLE_RULE.match(lines[i]):
                    block.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            head, rows = [plain(c) for c in block[0]], [[plain(c) for c in r] for r in block[1:]]
            t = doc.table(head, rows, widths=widths(block), zebra=len(rows) > 6, first_col_bold=True)
            keep_table_together(t, len(rows))
            continue
        elif re.match(r"\d+\. ", line):
            items = []
            while i < len(lines) and re.match(r"\d+\. ", lines[i]):
                items.append(re.sub(r"^\d+\. ", "", lines[i].rstrip()))
                i += 1
            numbered(doc, items)
            continue
        elif re.match(r"[-*] ", line):
            items = []
            while i < len(lines) and re.match(r"[-*] ", lines[i]):
                items.append(lines[i][2:].rstrip())
                i += 1
            bullets(doc, items)
            continue
        else:
            par = doc.rich(inline(line))
            if TABLE_ROW.match(next_line(lines, i)):         # บรรทัดนำตารางอยู่หน้าเดียวกับตาราง
                keep_with_next(par)
        i += 1

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


WD_EXPORT_FORMAT_PDF = 17   # ค่า wdExportFormatPDF ของ Word


# -- PDF สำหรับตรวจด้วยตา ----------------------------------------------------------------------------
def to_pdf(docx: Path) -> Path:
    """Word บน Windows (อัปเดตสารบัญให้ด้วย) ถ้าไม่มี → LibreOffice"""
    pdf = docx.with_suffix(".pdf")
    # path ส่งผ่าน environment variable ไม่ต่อเข้า string คำสั่ง → ชื่อไฟล์ที่มี ' ฉีดคำสั่งไม่ได้
    ps = ("$w=New-Object -ComObject Word.Application;$w.Visible=$false;"
          "$d=$w.Documents.Open($env:MD2DOCX_IN);$d.TablesOfContents | % {$_.Update()};$d.Save();"
          f"$d.ExportAsFixedFormat($env:MD2DOCX_OUT,{WD_EXPORT_FORMAT_PDF});$d.Close();$w.Quit()")
    env = {**os.environ, "MD2DOCX_IN": str(docx.resolve()), "MD2DOCX_OUT": str(pdf.resolve())}
    try:
        subprocess.run(["powershell", "-NoProfile", "-Command", ps], env=env, check=True, capture_output=True, timeout=300)
    except (OSError, subprocess.SubprocessError):
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(docx.parent), str(docx)],
                       check=True, capture_output=True, timeout=300)
    return pdf


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("source", type=Path, help="ไฟล์ Markdown ต้นฉบับ")
    ap.add_argument("-o", "--out", type=Path, help="ไฟล์ .docx ที่จะเขียน (ค่าเริ่มต้น: ชื่อเดียวกับต้นฉบับ)")
    ap.add_argument("--pdf", action="store_true", help="สร้าง .pdf ไว้ตรวจหน้าตาด้วย")
    a = ap.parse_args()
    out = build(a.source, a.out or a.source.with_suffix(".docx"))
    print("wrote", out)
    if a.pdf:
        print("wrote", to_pdf(out))


if __name__ == "__main__":
    main()
